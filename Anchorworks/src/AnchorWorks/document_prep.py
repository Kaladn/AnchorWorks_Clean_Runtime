from __future__ import annotations

import csv
import email
from email import policy
import hashlib
import html.parser
import io
import json
import posixpath
import re
import wave
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from .document_film import extract_pdf_document_film_from_bytes
from .visual_manifest import manifest_from_image_bytes
from .visual_recognition_layer import create_empty_recognition_layer
from .visual_region_map import create_empty_region_map


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".log",
    ".text",
    ".csv.txt",
    ".py",
    ".js",
    ".ts",
    ".css",
    ".scss",
    ".jsonl",
    ".yaml",
    ".yml",
    ".ini",
    ".cfg",
    ".toml",
}
DELIMITED_EXTENSIONS = {".csv", ".tsv", ".psv"}
JSON_EXTENSIONS = {".json", ".jsonld", ".geojson"}
XML_EXTENSIONS = {".xml", ".xhtml", ".opf", ".svg", ".cnxml"}
HTML_EXTENSIONS = {".html", ".htm"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xlsm"}
DOCX_EXTENSIONS = {".docx"}
PDF_EXTENSIONS = {".pdf"}
RTF_EXTENSIONS = {".rtf"}
ODT_EXTENSIONS = {".odt"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}
EPUB_EXTENSIONS = {".epub"}
PPTX_EXTENSIONS = {".pptx"}
EMAIL_EXTENSIONS = {".eml"}
ARCHIVE_EXTENSIONS = {".zip"}
VIDEO_EXTENSIONS = {".mp4", ".m4v", ".mov", ".avi", ".mkv", ".webm"}
AUDIO_EXTENSIONS = {".wav", ".mp3", ".m4a", ".flac", ".ogg"}
VISUAL_LEADER_RE = re.compile(r"[\u2500-\u257F]+")
NO_WRITE_POLICY = {"maps": False, "counts": False, "lifetime": False, "lexicon": False}
BLOCKED_ARCHIVE_EXTENSIONS = {".exe", ".dll", ".bat", ".cmd", ".ps1", ".msi", ".com", ".scr"}


@dataclass
class PreparedDocument:
    source_name: str
    source_path: str
    file_type: str
    original_size: int
    sha256: str
    converter: str
    prepared_text: str
    warnings: list[str] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def prepare_file(source_path: Path) -> PreparedDocument:
    path = Path(source_path).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(path)
    if path.is_dir():
        raise IsADirectoryError(path)
    return prepare_bytes(
        path.read_bytes(),
        source_name=path.name,
        source_path=str(path),
        file_type=_guess_file_type(path.name),
    )


def prepare_bytes(
    raw: bytes,
    *,
    source_name: str,
    source_path: str = "",
    file_type: str = "",
) -> PreparedDocument:
    source_name = Path(source_name or "document").name
    suffix = Path(source_name).suffix.lower()
    file_type = file_type or _guess_file_type(source_name)
    warnings: list[str] = []
    metadata: dict[str, Any] = {}

    try:
        if suffix in DELIMITED_EXTENSIONS:
            text, metadata = _prepare_delimited(raw, source_name, suffix)
            converter = "delimited-table"
        elif suffix in JSON_EXTENSIONS:
            text, metadata = _prepare_json(raw, source_name)
            converter = str(metadata.pop("_converter", "json-structure"))
        elif suffix in HTML_EXTENSIONS:
            text, metadata = _prepare_html(raw, source_name)
            converter = "html-text"
        elif suffix in XML_EXTENSIONS:
            text, metadata = _prepare_xml(raw, source_name)
            converter = "xml-structure"
        elif suffix in PDF_EXTENSIONS:
            text, metadata = _prepare_pdf(raw, source_name)
            converter = "pdf-text"
        elif suffix in DOCX_EXTENSIONS:
            text, metadata = _prepare_docx(raw, source_name)
            converter = "docx-text"
        elif suffix in SPREADSHEET_EXTENSIONS:
            text, metadata = _prepare_spreadsheet(raw, source_name)
            converter = "spreadsheet-table"
        elif suffix in RTF_EXTENSIONS:
            text, metadata = _prepare_rtf(raw, source_name)
            converter = "rtf-text"
        elif suffix in ODT_EXTENSIONS:
            text, metadata = _prepare_odt(raw, source_name)
            converter = "odt-text"
        elif suffix in IMAGE_EXTENSIONS:
            text, metadata = _prepare_image(raw, source_name)
            converter = "image-metadata"
        elif suffix in EPUB_EXTENSIONS:
            text, metadata = _prepare_epub(raw, source_name)
            converter = "epub-spine"
        elif suffix in PPTX_EXTENSIONS:
            text, metadata = _prepare_pptx(raw, source_name)
            converter = "pptx-text"
        elif suffix in EMAIL_EXTENSIONS:
            text, metadata = _prepare_email(raw, source_name)
            converter = "email-message"
        elif suffix in ARCHIVE_EXTENSIONS:
            text, metadata = _prepare_archive(raw, source_name)
            converter = "archive-zip"
        elif suffix in VIDEO_EXTENSIONS:
            text, metadata = _prepare_video(raw, source_name)
            converter = "video-evidence-manifest"
        elif suffix in AUDIO_EXTENSIONS:
            text, metadata = _prepare_audio(raw, source_name)
            converter = "audio-evidence-manifest"
        else:
            text = _decode_text(raw)
            metadata = {"text_characters": len(text)}
            converter = "plain-text"
    except Exception as exc:
        try:
            text = _decode_text(raw)
            converter = "plain-text-fallback"
            warnings.append(f"{suffix or 'unknown'} converter failed: {exc}")
            metadata = {"text_characters": len(text)}
        except Exception:
            text = _prepare_binary_fallback(raw, source_name)
            converter = "binary-metadata-fallback"
            warnings.append(f"{suffix or 'unknown'} converter failed: {exc}")
            metadata = {"binary_fallback": True}

    text, stripped_leader_count = _strip_visual_leaders(text)
    if stripped_leader_count:
        metadata["stripped_visual_leader_characters"] = stripped_leader_count

    return PreparedDocument(
        source_name=source_name,
        source_path=source_path,
        file_type=file_type,
        original_size=len(raw),
        sha256=hashlib.sha256(raw).hexdigest(),
        converter=converter,
        prepared_text=text,
        warnings=warnings,
        metadata=metadata,
    )


def _decode_text(raw: bytes) -> str:
    for encoding in ["utf-8-sig", "utf-8", "utf-16", "utf-16-le", "utf-16-be", "cp1252"]:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("unable to decode file as text")


def _strip_visual_leaders(text: str) -> tuple[str, int]:
    matches = list(VISUAL_LEADER_RE.finditer(text))
    if not matches:
        return text, 0
    stripped_count = sum(match.end() - match.start() for match in matches)
    return VISUAL_LEADER_RE.sub(" ", text), stripped_count


def _repair_common_mojibake(text: str) -> str:
    replacements = {
        "â€™": "’",
        "â€˜": "‘",
        "â€œ": "“",
        "â€\u009d": "”",
        "â€": "”",
        "â€“": "–",
        "â€”": "—",
        "â€¦": "…",
        "â€²": "′",
        "âˆ’": "−",
        "âˆž": "∞",
        "âˆ‘": "∑",
        "âˆª": "∪",
        "âˆ©": "∩",
        "âˆ˜": "∘",
        "â‹…": "⋅",
        "âŸ¶": "⟶",
        "â‡Œ": "⇌",
        "â‰¤": "≤",
        "â‰¥": "≥",
        "â‰ˆ": "≈",
        "Ã—": "×",
        "Ã·": "÷",
        "Â·": "·",
        "Â®": "®",
        "Â¯": "¯",
        "Â°": "°",
        "Â±": "±",
        "Âµ": "µ",
        "Â’": "’",
        "Ï€": "π",
        "Ïƒ": "σ",
        "Ï‰": "ω",
        "Ïˆ": "ψ",
        "Ï‡": "χ",
        "Î±": "α",
        "Î²": "β",
        "Î³": "γ",
        "Î´": "δ",
        "Î¸": "θ",
        "Î»": "λ",
        "Î¼": "μ",
    }
    repaired = str(text or "")
    for bad, good in replacements.items():
        repaired = repaired.replace(bad, good)
    if any(marker in repaired for marker in ("Ã", "Â", "Î", "Ï", "Å", "Ê", "ï¬")):
        try:
            decoded = repaired.encode("latin-1").decode("utf-8")
            repaired = decoded
        except UnicodeError:
            pass
    repaired = repaired.replace("ﬁ", "fi").replace("ﬂ", "fl")
    return repaired


def _prepare_delimited(raw: bytes, source_name: str, suffix: str) -> tuple[str, dict[str, Any]]:
    text = _decode_text(raw)
    delimiter = "\t" if suffix == ".tsv" else "|" if suffix == ".psv" else ","
    if suffix == ".csv":
        try:
            dialect = csv.Sniffer().sniff(text[:8192])
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","

    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    lines = [_source_line(source_name), "[TYPE: table]", "", f"[TABLE: {Path(source_name).stem}]"]
    lines.extend(" | ".join(_clean_cell(cell) for cell in row) for row in rows)
    lines.append("[TABLE_END]")
    return "\n".join(lines), {"rows": len(rows), "delimiter": delimiter}


def _prepare_json(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    data = json.loads(_decode_text(raw))
    rows: list[str] = []
    _flatten_json(data, "root", rows)
    lines = [_source_line(source_name), "[TYPE: json]", "", f"[JSON: {Path(source_name).stem}]"]
    lines.extend(rows or ["root | null |"])
    lines.append("[JSON_END]")
    return "\n".join(lines), {"rows": len(rows)}


def _flatten_json(value: Any, path: str, rows: list[str]) -> None:
    if isinstance(value, dict):
        if not value:
            rows.append(f"{path} | object | empty")
        for key, child in value.items():
            safe_key = str(key).replace("\n", " ").strip()
            _flatten_json(child, f"{path}.{safe_key}", rows)
    elif isinstance(value, list):
        if not value:
            rows.append(f"{path} | array | empty")
        for index, child in enumerate(value):
            _flatten_json(child, f"{path}.{index}", rows)
    else:
        rows.append(f"{path} | {type(value).__name__} | {_clean_cell(value)}")


def _prepare_xml(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    root = ElementTree.fromstring(raw)
    rows: list[str] = []
    _walk_xml(root, rows)
    return "\n".join(rows), {"rows": len(rows), "root": _strip_namespace(root.tag)}


def _walk_xml(element: ElementTree.Element, rows: list[str]) -> None:
    text = (element.text or "").strip()
    if text:
        rows.append(_repair_common_mojibake(_clean_cell(text)))
    for child in list(element):
        _walk_xml(child, rows)
        tail = (child.tail or "").strip()
        if tail:
            rows.append(_repair_common_mojibake(_clean_cell(tail)))


def _prepare_html(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    parser = _HTMLTextExtractor()
    parser.feed(_decode_text(raw))
    parser.close()
    lines = [_source_line(source_name), "[TYPE: html]", "", f"[HTML: {Path(source_name).stem}]"]
    lines.extend(parser.lines)
    lines.append("[HTML_END]")
    return "\n".join(lines), {
        "lines": len(parser.lines),
        "attribute_sidecars": parser.attribute_sidecars,
        "visual_refs": parser.visual_refs,
        "writes_allowed": dict(NO_WRITE_POLICY),
    }


def _prepare_pdf(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except Exception as exc:  # pragma: no cover - depends on optional runtime packages
        raise ValueError("PDF conversion needs pypdf installed") from exc

    reader = PdfReader(io.BytesIO(raw))
    lines = [_source_line(source_name), "[TYPE: pdf]", "", f"[PDF: {Path(source_name).stem}]"]
    page_locators: list[dict[str, Any]] = []
    scanned_pages: list[int] = []
    for page_index, page in enumerate(reader.pages, start=1):
        extracted = page.extract_text() or ""
        lines.append("")
        lines.append(f"[PAGE: {page_index}]")
        lines.append(extracted)
        page_locators.append({"page": page_index, "line_start": len(lines), "text_present": bool(extracted.strip())})
        if not extracted.strip():
            scanned_pages.append(page_index)
    lines.append("[PDF_END]")
    metadata: dict[str, Any] = {"pages": len(reader.pages), "page_locators": page_locators, "writes_allowed": dict(NO_WRITE_POLICY)}
    if scanned_pages:
        metadata["scanned_or_image_only_pages"] = scanned_pages
        metadata["warnings"] = ["PDF page text was empty; OCR is not run by this converter."]
    try:
        metadata["document_film"] = extract_pdf_document_film_from_bytes(
            raw,
            source_name=source_name,
            source_path=source_name,
            frame_rate=1.0,
            duplicate_policy="all_pages_report_duplicates",
        )
    except ValueError as exc:
        metadata["document_film_status"] = "not_available"
        metadata["document_film_warning"] = str(exc)
    return "\n".join(lines), metadata


def _prepare_docx(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover - depends on optional runtime packages
        raise ValueError("DOCX conversion needs python-docx installed") from exc

    doc = Document(io.BytesIO(raw))
    lines = [_source_line(source_name), "[TYPE: docx]", "", f"[DOCX: {Path(source_name).stem}]"]
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for table_index, table in enumerate(doc.tables, start=1):
        lines.extend(["", f"[TABLE: {Path(source_name).stem} {table_index}]"])
        for row in table.rows:
            row_items = getattr(row, "cells")
            lines.append(" | ".join(_clean_cell(table_item.text) for table_item in row_items))
        lines.append("[TABLE_END]")
    lines.append("[DOCX_END]")
    sidecars = _inspect_docx_sidecars(raw)
    return "\n".join(lines), {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        **sidecars,
        "writes_allowed": dict(NO_WRITE_POLICY),
    }


def _prepare_spreadsheet(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - depends on optional runtime packages
        raise ValueError("spreadsheet conversion needs openpyxl installed") from exc

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    lines = [_source_line(source_name), "[TYPE: spreadsheet]", "", f"[SPREADSHEET: {Path(source_name).stem}]"]
    row_count = 0
    for sheet in workbook.worksheets:
        lines.extend(["", f"[TABLE: {sheet.title}]"])
        for row in sheet.iter_rows(values_only=True):
            values = [_clean_cell(value) for value in row]
            if any(values):
                row_count += 1
                lines.append(" | ".join(values))
        lines.append("[TABLE_END]")
    lines.append("[SPREADSHEET_END]")
    return "\n".join(lines), {"sheets": len(workbook.worksheets), "rows": row_count}


def _prepare_rtf(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    text = _decode_text(raw)
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
    text = text.replace("{", " ").replace("}", " ").replace("\\", " ")
    text = re.sub(r"[ \t]+", " ", text)
    lines = [_source_line(source_name), "[TYPE: rtf]", "", f"[RTF: {Path(source_name).stem}]", text.strip(), "[RTF_END]"]
    return "\n".join(lines), {"text_characters": len(text)}


def _prepare_odt(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        content = archive.read("content.xml")
    root = ElementTree.fromstring(content)
    rows: list[str] = []
    for element in root.iter():
        text = "".join(element.itertext()).strip()
        if text:
            rows.append(_clean_cell(text))
    lines = [_source_line(source_name), "[TYPE: odt]", "", f"[ODT: {Path(source_name).stem}]"]
    lines.extend(dict.fromkeys(rows))
    lines.append("[ODT_END]")
    return "\n".join(lines), {"lines": len(rows)}


def _prepare_image(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    visual_manifest = manifest_from_image_bytes(raw, source_name)
    visual_region_map = create_empty_region_map(visual_manifest)
    visual_recognition_layer = create_empty_recognition_layer(visual_manifest, visual_region_map)
    source = visual_manifest.source
    metadata: dict[str, Any] = {
        "format": source.file_format,
        "width": source.width,
        "height": source.height,
        "mode": source.color_mode,
        "visual_manifest": visual_manifest.to_dict(),
        "visual_region_map": visual_region_map.to_dict(),
        "visual_recognition_layer": visual_recognition_layer.to_dict(),
        "visual_authority": visual_manifest.authority,
        "visual_approval_status": visual_manifest.approval_status,
    }
    warnings: list[str] = []
    if source.width is None or source.height is None:
        warnings.append("image geometry could not be read")

    lines = [
        _source_line(source_name),
        "[TYPE: image]",
        "",
        f"[IMAGE: {Path(source_name).stem}]",
        f"Visual_Record_ID: {source.visual_record_id}",
        f"Format: {metadata.get('format', Path(source_name).suffix.lower().lstrip('.'))}",
        f"Width: {metadata.get('width', 'unknown')}",
        f"Height: {metadata.get('height', 'unknown')}",
        f"Aspect_Ratio: {source.aspect_ratio}",
        f"Mode: {metadata.get('mode', 'unknown')}",
        f"SHA256: {hashlib.sha256(raw).hexdigest()}",
        f"Region_Map_ID: {visual_region_map.region_map_id}",
        f"Recognition_Layer_ID: {visual_recognition_layer.recognition_layer_id}",
        "Authority: source_local_visual_evidence",
        "Approval_Status: preview_only",
        "Writes_Allowed: maps=false counts=false lifetime=false lexicon=false",
        "Visual_Note: Native pixels are source evidence. OCR, object, and scene layers are derived later.",
        "[IMAGE_END]",
    ]
    if warnings:
        metadata["warnings"] = warnings
    return "\n".join(lines), metadata


def _prepare_epub(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        rootfile = _epub_rootfile(archive)
        root_dir = posixpath.dirname(rootfile)
        opf_root = ElementTree.fromstring(archive.read(rootfile))
        manifest: dict[str, dict[str, str]] = {}
        for item in opf_root.iter():
            if _strip_namespace(item.tag) != "item":
                continue
            item_id = item.attrib.get("id")
            href = item.attrib.get("href")
            if item_id and href:
                full_path = posixpath.normpath(posixpath.join(root_dir, href))
                manifest[item_id] = {
                    "href": href,
                    "full_path": full_path,
                    "media_type": item.attrib.get("media-type", ""),
                }
        spine_ids = [
            item.attrib.get("idref", "")
            for item in opf_root.iter()
            if _strip_namespace(item.tag) == "itemref" and item.attrib.get("idref")
        ]
        lines = [_source_line(source_name), "[TYPE: epub]", "", f"[EPUB: {Path(source_name).stem}]"]
        visual_refs: list[dict[str, Any]] = []
        spine_count = 0
        for spine_index, item_id in enumerate(spine_ids, start=1):
            item = manifest.get(item_id)
            if not item:
                continue
            chapter_path = item["full_path"]
            if chapter_path not in archive.namelist():
                continue
            spine_count += 1
            html_text, html_meta = _prepare_html(archive.read(chapter_path), Path(chapter_path).name)
            lines.extend(["", f"[EPUB_SPINE_ITEM: {spine_index} {chapter_path}]"])
            lines.extend(_body_lines(html_text))
            for ref in html_meta.get("visual_refs") or []:
                source_ref = str(ref.get("source_path") or "")
                if source_ref:
                    ref = dict(ref)
                    ref["source_path"] = posixpath.normpath(posixpath.join(posixpath.dirname(chapter_path), source_ref))
                    ref["spine_index"] = spine_index
                    ref["chapter_path"] = chapter_path
                    visual_refs.append(ref)
        lines.append("[EPUB_END]")
    return "\n".join(lines), {
        "rootfile": rootfile,
        "spine_count": spine_count,
        "manifest_count": len(manifest),
        "visual_refs": visual_refs,
        "writes_allowed": dict(NO_WRITE_POLICY),
    }


def _epub_rootfile(archive: zipfile.ZipFile) -> str:
    try:
        container = ElementTree.fromstring(archive.read("META-INF/container.xml"))
    except KeyError as exc:
        raise ValueError("EPUB missing META-INF/container.xml") from exc
    for element in container.iter():
        if _strip_namespace(element.tag) == "rootfile" and element.attrib.get("full-path"):
            return element.attrib["full-path"]
    raise ValueError("EPUB container missing rootfile")


def _prepare_pptx(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    lines = [_source_line(source_name), "[TYPE: pptx]", "", f"[PPTX: {Path(source_name).stem}]"]
    embedded_media: list[dict[str, Any]] = []
    slide_count = 0
    note_count = 0
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        names = set(archive.namelist())
        for slide_path in sorted((name for name in names if re.match(r"ppt/slides/slide\d+\.xml$", name)), key=_natural_sort_key):
            slide_count += 1
            root = ElementTree.fromstring(archive.read(slide_path))
            lines.extend(["", f"[SLIDE: {slide_count}]"])
            for text in _xml_text_values(root, "t"):
                lines.append(_clean_cell(text))
            for pic_name in _pptx_picture_names(root):
                embedded_media.append({
                    "kind": "embedded_image_placeholder",
                    "source_path": f"{slide_path}#{pic_name}",
                    "slide": slide_count,
                    "recognition_status": "not_run",
                    "writes_allowed": dict(NO_WRITE_POLICY),
                })
            notes_path = f"ppt/notesSlides/notesSlide{slide_count}.xml"
            if notes_path in names:
                note_count += 1
                notes_root = ElementTree.fromstring(archive.read(notes_path))
                lines.append(f"[NOTES: {slide_count}]")
                for text in _xml_text_values(notes_root, "t"):
                    lines.append(_clean_cell(text))
    lines.append("[PPTX_END]")
    return "\n".join(lines), {
        "slides": slide_count,
        "notes": note_count,
        "embedded_media": embedded_media,
        "writes_allowed": dict(NO_WRITE_POLICY),
    }


def _prepare_email(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    message = email.message_from_bytes(raw, policy=policy.default)
    lines = [_source_line(source_name), "[TYPE: email]", "", f"[EMAIL: {Path(source_name).stem}]"]
    for field in ("Subject", "From", "To", "Date"):
        value = message.get(field)
        if value:
            lines.append(f"{field}: {value}")
    body_lines: list[str] = []
    attachments: list[dict[str, Any]] = []
    for part in message.walk():
        if part.is_multipart():
            continue
        filename = part.get_filename()
        payload = part.get_payload(decode=True) or b""
        content_type = part.get_content_type()
        if filename:
            child = prepare_bytes(payload, source_name=filename, source_path=f"{source_name}!/{filename}", file_type=content_type)
            child_record = {
                "filename": filename,
                "content_type": content_type,
                "size_bytes": len(payload),
                "converter": child.converter,
                "warnings": child.warnings,
                "writes_allowed": dict(NO_WRITE_POLICY),
            }
            if child.metadata.get("visual_manifest"):
                child_record["visual_manifest"] = child.metadata["visual_manifest"]
            attachments.append(child_record)
            continue
        if content_type == "text/plain":
            body_lines.append(_decode_text(payload))
        elif content_type == "text/html":
            html_text, _ = _prepare_html(payload, source_name)
            body_lines.extend(_body_lines(html_text))
    if body_lines:
        lines.extend(["", "[EMAIL_BODY]"])
        lines.extend(_clean_cell(line) for line in "\n".join(body_lines).splitlines() if _clean_cell(line))
        lines.append("[EMAIL_BODY_END]")
    lines.append("[EMAIL_END]")
    return "\n".join(lines), {
        "subject": str(message.get("Subject") or ""),
        "from": str(message.get("From") or ""),
        "to": str(message.get("To") or ""),
        "date": str(message.get("Date") or ""),
        "attachments": attachments,
        "writes_allowed": dict(NO_WRITE_POLICY),
    }


def _prepare_archive(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    lines = [_source_line(source_name), "[TYPE: archive]", "", f"[ARCHIVE: {Path(source_name).stem}]"]
    children: list[dict[str, Any]] = []
    blocked: list[dict[str, Any]] = []
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        for info in sorted(archive.infolist(), key=lambda item: item.filename.lower()):
            name = info.filename
            if info.is_dir():
                continue
            suffix = Path(name).suffix.lower()
            if suffix in BLOCKED_ARCHIVE_EXTENSIONS:
                blocked.append({"path": name, "reason": "blocked_executable_member", "size_bytes": info.file_size})
                continue
            if suffix in ARCHIVE_EXTENSIONS:
                blocked.append({"path": name, "reason": "nested_archive_not_expanded", "size_bytes": info.file_size})
                continue
            child_raw = archive.read(info)
            child = prepare_bytes(child_raw, source_name=Path(name).name, source_path=f"{source_name}!/{name}")
            children.append({
                "path": name,
                "converter": child.converter,
                "size_bytes": info.file_size,
                "warnings": child.warnings,
                "writes_allowed": dict(NO_WRITE_POLICY),
            })
            lines.extend(["", f"[ARCHIVE_CHILD: {name}]"])
            lines.extend(_body_lines(child.prepared_text))
    lines.append("[ARCHIVE_END]")
    return "\n".join(lines), {
        "children": children,
        "blocked_children": blocked,
        "writes_allowed": dict(NO_WRITE_POLICY),
    }


def _prepare_video(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    metadata = {
        "media_type": "video",
        "size_bytes": len(raw),
        "sha256": hashlib.sha256(raw).hexdigest(),
        "duration_seconds": None,
        "frame_rate": None,
        "width": None,
        "height": None,
        "frame_records": [],
        "recognition_status": "not_run",
        "writes_allowed": dict(NO_WRITE_POLICY),
    }
    lines = [
        _source_line(source_name),
        "[TYPE: video]",
        "",
        f"[VIDEO: {Path(source_name).stem}]",
        f"SHA256: {metadata['sha256']}",
        "Recognition_Status: not_run",
        "Writes_Allowed: maps=false counts=false lifetime=false lexicon=false",
        "[VIDEO_END]",
    ]
    return "\n".join(lines), metadata


def _prepare_audio(raw: bytes, source_name: str) -> tuple[str, dict[str, Any]]:
    metadata: dict[str, Any] = {
        "media_type": "audio",
        "size_bytes": len(raw),
        "sha256": hashlib.sha256(raw).hexdigest(),
        "duration_seconds": None,
        "sample_rate": None,
        "channels": None,
        "transcript_status": "not_supplied",
        "writes_allowed": dict(NO_WRITE_POLICY),
    }
    if Path(source_name).suffix.lower() == ".wav":
        with wave.open(io.BytesIO(raw), "rb") as wav:
            frames = wav.getnframes()
            rate = wav.getframerate()
            metadata["sample_rate"] = rate
            metadata["channels"] = wav.getnchannels()
            metadata["duration_seconds"] = frames / rate if rate else None
    lines = [
        _source_line(source_name),
        "[TYPE: audio]",
        "",
        f"[AUDIO: {Path(source_name).stem}]",
        f"SHA256: {metadata['sha256']}",
        f"Sample_Rate: {metadata.get('sample_rate') or 'unknown'}",
        "Transcript_Status: not_supplied",
        "Writes_Allowed: maps=false counts=false lifetime=false lexicon=false",
        "[AUDIO_END]",
    ]
    return "\n".join(lines), metadata


def _inspect_docx_sidecars(raw: bytes) -> dict[str, Any]:
    sidecars: dict[str, Any] = {
        "headers": [],
        "footers": [],
        "footnotes": [],
        "endnotes": [],
        "comments": [],
        "embedded_media": [],
        "warnings": [],
    }
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            for name in archive.namelist():
                lower = name.lower()
                if lower.startswith("word/media/"):
                    sidecars["embedded_media"].append({
                        "kind": "embedded_media_placeholder",
                        "source_path": name,
                        "recognition_status": "not_run",
                        "writes_allowed": dict(NO_WRITE_POLICY),
                    })
                elif lower.startswith("word/header") and lower.endswith(".xml"):
                    sidecars["headers"].extend(_xml_file_text(archive, name))
                elif lower.startswith("word/footer") and lower.endswith(".xml"):
                    sidecars["footers"].extend(_xml_file_text(archive, name))
                elif lower == "word/footnotes.xml":
                    sidecars["footnotes"].extend(_xml_file_text(archive, name))
                elif lower == "word/endnotes.xml":
                    sidecars["endnotes"].extend(_xml_file_text(archive, name))
                elif lower == "word/comments.xml":
                    sidecars["comments"].extend(_xml_file_text(archive, name))
                elif lower == "word/document.xml":
                    text = archive.read(name).decode("utf-8", errors="ignore")
                    if "<w:ins" in text or "<w:del" in text:
                        sidecars["warnings"].append("DOCX contains tracked-change markup; changes are not resolved by this converter.")
    except Exception:
        pass
    return sidecars


def _prepare_binary_fallback(raw: bytes, source_name: str) -> str:
    sample = raw[:2048].hex(" ")
    return "\n".join(
        [
            _source_line(source_name),
            "[TYPE: binary]",
            "",
            f"[BINARY_FILE: {Path(source_name).stem}]",
            f"Extension: {Path(source_name).suffix.lower() or 'none'}",
            f"Size_Bytes: {len(raw)}",
            f"SHA256: {hashlib.sha256(raw).hexdigest()}",
            "[BYTE_HEX_SAMPLE_START]",
            sample,
            "[BYTE_HEX_SAMPLE_END]",
        ]
    )


class _HTMLTextExtractor(html.parser.HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.lines: list[str] = []
        self.attribute_sidecars: list[dict[str, Any]] = []
        self.visual_refs: list[dict[str, Any]] = []
        self._current: list[str] = []
        self._skip_depth = 0
        self._cell_row: list[str] | None = None
        self._cell_text: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        attr_map = {name.lower(): value for name, value in attrs if name and value}
        for attr in ("href", "alt", "title", "aria-label"):
            if attr in attr_map:
                self.attribute_sidecars.append({
                    "tag": tag,
                    "attribute": attr,
                    "value": attr_map[attr],
                    "authority": "source_local_html_attribute",
                    "writes_allowed": dict(NO_WRITE_POLICY),
                })
        if tag == "img":
            self.visual_refs.append({
                "visual_record_id": "",
                "kind": "image_reference",
                "source_path": attr_map.get("src", ""),
                "alt_text": attr_map.get("alt", ""),
                "title": attr_map.get("title", ""),
                "geometry_status": "unknown",
                "recognition_status": "not_run",
                "writes_allowed": dict(NO_WRITE_POLICY),
            })
        if tag in {"script", "style", "noscript"}:
            self._skip_depth += 1
            return
        if tag == "tr":
            self._flush_line()
            self._cell_row = []
        elif tag in {"td", "th"}:
            self._cell_text = []
        elif tag in {"br", "p", "div", "section", "article", "header", "footer", "h1", "h2", "h3", "h4", "h5", "h6", "li"}:
            self._flush_line()

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"script", "style", "noscript"}:
            self._skip_depth = max(0, self._skip_depth - 1)
            return
        if tag in {"td", "th"} and self._cell_row is not None:
            self._cell_row.append(_clean_cell(" ".join(self._cell_text or [])))
            self._cell_text = None
        elif tag == "tr" and self._cell_row is not None:
            if any(self._cell_row):
                self.lines.append(" | ".join(self._cell_row))
            self._cell_row = None
        elif tag in {"p", "div", "section", "article", "header", "footer", "h1", "h2", "h3", "h4", "h5", "h6", "li"}:
            self._flush_line()

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        text = data.strip()
        if not text:
            return
        if self._cell_text is not None:
            self._cell_text.append(text)
        else:
            self._current.append(text)

    def close(self) -> None:
        self._flush_line()
        super().close()

    def _flush_line(self) -> None:
        if self._current:
            line = _clean_cell(" ".join(self._current))
            if line:
                self.lines.append(line)
            self._current = []


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _source_line(source_name: str) -> str:
    return f"[SOURCE: {source_name}]"


def _strip_namespace(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _guess_file_type(source_name: str) -> str:
    suffix = Path(source_name).suffix.lower().lstrip(".")
    return suffix or "unknown"


def _body_lines(prepared_text: str) -> list[str]:
    lines = []
    for line in prepared_text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("[SOURCE:") or stripped.startswith("[TYPE:"):
            continue
        lines.append(stripped)
    return lines


def _natural_sort_key(value: str) -> list[Any]:
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", value)]


def _xml_text_values(root: ElementTree.Element, local_name: str) -> list[str]:
    values: list[str] = []
    for element in root.iter():
        if _strip_namespace(element.tag) == local_name:
            text = "".join(element.itertext()).strip()
            if text:
                values.append(text)
    return values


def _xml_file_text(archive: zipfile.ZipFile, name: str) -> list[str]:
    try:
        root = ElementTree.fromstring(archive.read(name))
    except Exception:
        return []
    return [_clean_cell(text) for text in _xml_text_values(root, "t") if _clean_cell(text)]


def _pptx_picture_names(root: ElementTree.Element) -> list[str]:
    names: list[str] = []
    for element in root.iter():
        if _strip_namespace(element.tag) == "cNvPr":
            name = element.attrib.get("name", "")
            if "picture" in name.lower() or "image" in name.lower():
                names.append(name)
    return names
